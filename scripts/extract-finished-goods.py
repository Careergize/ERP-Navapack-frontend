"""Extract the supplied workbook into frontend reference data; never edits Excel.
Usage: python scripts/extract-finished-goods.py <workbook.xlsx>
Requires openpyxl for this development-time task only.
"""
import datetime
import hashlib
import json
import pathlib
import sys
from collections import Counter
import openpyxl
from openpyxl.utils import get_column_letter

path = pathlib.Path(sys.argv[1])
source = openpyxl.load_workbook(path, data_only=False)
cached = openpyxl.load_workbook(path, data_only=True)
products, movements = [], []
categories = {'Carrier Bags': 'carrier-bags', 'Flat Bags': 'flat-bags', 'Packing': 'packing'}

def text(value):
    return ' '.join(str(value).split()) if value is not None else ''

def number(sheet, row, col):
    value = sheet.cell(row, col).value
    if value is None or text(value) == '':
        return None
    # Read cached results for formulas, including arithmetic daily entries.
    assert not isinstance(value, bool)
    try:
        return round(float(value), 6)
    except (ValueError, TypeError):
        raise ValueError(f'Non-numeric stock value: {sheet.title}!{sheet.cell(row, col).coordinate}: {value}')

for sheet in source:
    values = cached[sheet.title]
    category = categories[sheet.title]
    header = 1 if category == 'packing' else 2
    subgroup = sheet.title
    customer_group = ''
    previous_item = ''
    for row in range(header + 2, sheet.max_row + 1):
        fields = [text(values.cell(row, col).value) for col in range(1, 8)]
        customer, item, size, brand, packing, _, raw_unit = fields
        # Merged subgroup labels are not stock variants. Skip numeric/formula-only tails.
        is_heading = any(r.min_row == row and r.max_row == row and r.max_col >= 5 and r.min_col <= 2 for r in sheet.merged_cells.ranges)
        if is_heading:
            subgroup = customer or item
            customer_group = previous_item = ''
            continue
        if not any([item, size, brand, packing]):
            continue
        notes = []
        source_customer = customer
        # Carrier blanks interleave generic catalogue rows with custom orders: do not
        # forward-fill them. Flat bread wrappers have explicit customer blocks.
        inherited = False
        if category == 'flat-bags':
            if customer:
                customer_group = customer
            elif subgroup == 'LD PRINTED BREAD WRAPPERS':
                customer = customer_group
                inherited = bool(customer)
                if not item:
                    item = previous_item
            elif row in {63, 68, 139, 159, 160}:
                # Inspected coffee colour variants, RK size continuation, and the
                # AK MD printed-roll block. Other blanks remain general stock.
                customer = customer_group
                inherited = bool(customer)
        previous_item = item
        unit = raw_unit.lower()
        if unit in {'kgs', 'kg'}:
            unit = 'kg'
        elif unit == 'ctn':
            unit = 'ctn'
        elif not unit:
            # Exactly two blank-unit variants, resolved from adjacent matching rows.
            assert (sheet.title, row) in {('Carrier Bags', 7), ('Flat Bags', 101)}
            unit = 'ctn' if category == 'carrier-bags' else 'kg'
            notes.append(f'Unit inferred as {unit} from matching neighbouring variants; source unit is blank.')
        else:
            raise ValueError(f'Unexpected unit: {raw_unit}')
        size = size.replace('x', 'X')
        packing = packing.upper()
        opening = number(values, row, 6)
        daily = []
        for col in range(8, 70):
            quantity = number(values, row, col)
            if quantity is None or quantity == 0:
                continue
            date = values.cell(header, col if col % 2 == 0 else col - 1).value
            assert isinstance(date, datetime.datetime)
            assert quantity > 0
            daily.append((col, date.strftime('%Y-%m-%d'), 'IN' if col % 2 == 0 else 'OUT', quantity))
        # Carrier/Flat formulas treat a blank opening as zero. Packing is an empty
        # template: blank quantities are unrecorded, not evidence of zero inventory.
        if opening is None and category != 'packing':
            opening = 0
        if category == 'packing':
            notes.append('Packing sheet has May 2026 date headers and a conflicting opening-stock label (03/05/26). No opening quantity or movements were recorded.')
        stock_in = sum(q for _, _, direction, q in daily if direction == 'IN')
        stock_out = sum(q for _, _, direction, q in daily if direction == 'OUT')
        closing = None if opening is None else round(opening + stock_in - stock_out, 6)
        source_closing = number(values, row, 72)
        if closing is not None and source_closing is not None:
            assert abs(closing - source_closing) < 0.00001, (sheet.title, row, closing, source_closing)
        identity = '|'.join([category, subgroup, customer, item, size, brand, packing, unit])
        digest = hashlib.sha256(identity.casefold().encode()).hexdigest()[:12]
        product = dict(id=f'fg-{digest}', category=category, subgroup=subgroup, customer=customer,
                       itemDetails=item, size=size, brand=brand, packingSize=packing, unit=unit,
                       openingStock=opening, openingDate=values.cell(header, 8).value.strftime('%Y-%m-%d'),
                       sourceSheet=sheet.title, sourceRow=row, sourceCustomer=source_customer,
                       customerInherited=inherited, sourceUnit=raw_unit, sourceClosingStock=source_closing,
                       dataNotes=notes)
        products.append(product)
        product['_daily'] = daily

# Never collapse repeated rows: exact duplicate identities need client confirmation.
counts = Counter(p['id'] for p in products)
for product in products:
    if counts[product['id']] > 1:
        product['id'] += f"-row-{product['sourceRow']}"
        product['dataNotes'].append('Repeated variant identity in source; kept as a separate source row pending confirmation.')
    for col, date, direction, quantity in product.pop('_daily'):
        movements.append(dict(id=f"{product['id']}-{col}", productId=product['id'], date=date,
                              type=direction, quantity=quantity, unit=product['unit'],
                              reference=f"{product['sourceSheet']}!{get_column_letter(col)}{product['sourceRow']}",
                              customer=product['customer'], remarks='Client finished-goods workbook'))
target = pathlib.Path(__file__).resolve().parents[1] / 'src/lib/finishedGoodsReference.ts'
def array(data):
    return '[\n' + ',\n'.join('  ' + json.dumps(row, ensure_ascii=False) for row in data) + '\n]'
target.write_text('// Generated by scripts/extract-finished-goods.py. Source: ' + path.name + '\n'
                  + "import type { FinishedGood, FinishedGoodsMovement } from './finishedGoodsData';\n"
                  + 'export const FINISHED_GOODS_REFERENCE: FinishedGood[] = ' + array(products) + ';\n'
                  + 'export const FINISHED_MOVEMENTS_REFERENCE: FinishedGoodsMovement[] = ' + array(movements) + ';\n', encoding='utf-8')
print(json.dumps(dict(variants=dict(Counter(p['sourceSheet'] for p in products)),
                     movements=dict(Counter(m['date'][:7] for m in movements)),
                     totalVariants=len(products), duplicates=sum(n for n in counts.values() if n > 1),
                     inheritedCustomers=sum(p['customerInherited'] for p in products),
                     units=dict(Counter(p['unit'] for p in products))), indent=2))
